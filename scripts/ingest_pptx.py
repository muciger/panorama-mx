"""Ingesta desde el Reporte mensual de indicadores económicos (Presidencia INEGI).

Dos fuentes complementarias:
  1. PPT (python-pptx). Tabla reciente de 13 meses con columnas mensuales + anuales.
  2. XLSX datos embebidos (pandas). Serie histórica larga, generalmente anual.

El script mergea ambas según config/ppt_mapping.json. XLSX aporta profundidad histórica
(para MA12 con ventana completa). PPT sobrescribe últimos N periodos con columnas mensuales
y redondeos reportados.

Uso:
    python3 scripts/ingest_pptx.py <pptx> [--xlsx <xlsx>] [--indicator <id>] [--dry-run]

Ejemplos:
    python3 scripts/ingest_pptx.py sources/reporte.pptx --xlsx sources/datos.xlsx
    python3 scripts/ingest_pptx.py sources/reporte.pptx --xlsx sources/datos.xlsx --indicator inpc_mensual
    python3 scripts/ingest_pptx.py sources/reporte.pptx --dry-run
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from pptx import Presentation

MESES_ES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "octubre": 10,
    "noviembre": 11, "diciembre": 12,
}

MESES_ABR = {
    1: "ene", 2: "feb", 3: "mar", 4: "abr", 5: "may", 6: "jun",
    7: "jul", 8: "ago", 9: "sep", 10: "oct", 11: "nov", 12: "dic",
}

MESES_ABR_REV = {v: k for k, v in MESES_ABR.items()}

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
CONFIG_PATH = PROJECT_ROOT / "config" / "ppt_mapping.json"
DATA_DIR = PROJECT_ROOT / "data"


def to_num(s) -> float | int | None:
    """Convierte strings/valores numéricos. None si no parsea."""
    if s is None:
        return None
    if isinstance(s, (int, float)) and not isinstance(s, bool):
        if pd.isna(s):
            return None
        v = float(s)
    else:
        s = str(s).strip().replace(",", "")
        if s in ("", "-", "—", "n.d.", "N.D.", "ND", "nan"):
            return None
        try:
            v = float(s)
        except ValueError:
            return None
    if v.is_integer():
        return int(v)
    return round(v, 2)


def date_to_periodo(d) -> str | None:
    """datetime -> 'mmm-YY' en abreviaturas españolas."""
    if pd.isna(d) or d is None:
        return None
    if isinstance(d, str):
        return d.strip() or None
    try:
        month = int(d.month)
        year = int(d.year)
    except AttributeError:
        return None
    return f"{MESES_ABR[month]}-{year % 100:02d}"


def periodo_sort_key(periodo: str) -> int:
    """'mar-26' -> 2026*12 + 3 para sort cronológico. Retorna -1 si no parsea."""
    m = re.match(r"^([a-zñ]+)-(\d{2})$", periodo.lower().strip())
    if not m:
        return -1
    mes_abr, yy = m.group(1), int(m.group(2))
    month = MESES_ABR_REV.get(mes_abr)
    if month is None:
        return -1
    year = 2000 + yy
    return year * 12 + month


def compute_ma12(arr: list[float | None]) -> list[float | None]:
    """MA12 idéntica a normalize.compute_ma12. Ventana mínima: 3 observaciones."""
    out = []
    for i in range(len(arr)):
        window = [x for x in arr[max(0, i - 11):i + 1] if isinstance(x, (int, float))]
        out.append(round(sum(window) / len(window), 2) if len(window) >= 3 else None)
    return out


def parse_ppt_date(presentation: Presentation) -> str | None:
    """Fecha de portada -> 'YYYY-MM-DD'. e.g. '10 de abril de 2026'."""
    slide0 = presentation.slides[0]
    for shape in slide0.shapes:
        if not shape.has_text_frame:
            continue
        txt = shape.text_frame.text.strip().lower()
        m = re.search(r"(\d{1,2})\s+de\s+([a-záéíóú]+)\s+de\s+(\d{4})", txt)
        if m:
            day, mes, year = m.group(1), m.group(2), m.group(3)
            mes_num = MESES_ES.get(mes)
            if mes_num:
                return f"{year}-{mes_num:02d}-{int(day):02d}"
    return None


def get_slide_title(slide) -> str:
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if txt:
                return txt
    return ""


def find_slide(presentation: Presentation, rules: dict):
    title_contains = rules.get("title_contains", "")
    title_not_contains = rules.get("title_not_contains")
    idx = rules.get("slide_index")

    def title_ok(title: str) -> bool:
        if title_contains and title_contains not in title:
            return False
        if title_not_contains and title_not_contains in title:
            return False
        return True

    if idx is not None and idx < len(presentation.slides):
        slide = presentation.slides[idx]
        if title_ok(get_slide_title(slide)):
            return idx, slide

    for i, slide in enumerate(presentation.slides):
        if title_ok(get_slide_title(slide)):
            return i, slide

    return None, None


def extract_ppt_records(presentation: Presentation, cfg: dict) -> list[dict]:
    """Lee la tabla del slide y regresa lista de records con todas las columnas del column_map."""
    slide_idx, slide = find_slide(presentation, cfg)
    if slide is None:
        raise RuntimeError(f"no se encontró slide con title_contains='{cfg.get('title_contains')}'")

    table_index = cfg.get("table_index", 0)
    tables = [shape.table for shape in slide.shapes if shape.has_table]
    if table_index >= len(tables):
        raise RuntimeError(f"no hay tabla {table_index} en slide {slide_idx}")

    tbl = tables[table_index]
    rows = [[cell.text.strip() for cell in row.cells] for row in tbl.rows]
    header_rows = cfg.get("header_rows", 0)
    data_rows = rows[header_rows:]
    column_map = cfg["column_map"]

    records = []
    for row in data_rows:
        if len(row) < len(column_map):
            raise RuntimeError(f"fila con {len(row)} columnas, esperado {len(column_map)}: {row}")
        periodo = row[0].strip()
        if not periodo:
            continue
        record = {"Periodo": periodo}
        for i, col in enumerate(column_map[1:], start=1):
            record[col] = to_num(row[i])
        records.append(record)
    return records


def extract_xlsx_records(xlsx_path: Path, cfg: dict) -> list[dict]:
    """Lee una hoja del XLSX de datos embebidos y regresa records con las columnas mapeadas."""
    sheet = cfg["sheet"]
    header_row = cfg.get("header_row", 0)
    date_col = cfg.get("date_col", 0)
    column_map = cfg["column_map"]

    df = pd.read_excel(xlsx_path, sheet_name=sheet, header=None)
    headers = df.iloc[header_row].tolist()

    col_indexes = {}
    for xlsx_col, json_col in column_map.items():
        if xlsx_col not in headers:
            raise RuntimeError(f"columna '{xlsx_col}' no existe en hoja '{sheet}'. Headers: {headers}")
        col_indexes[json_col] = headers.index(xlsx_col)

    records = []
    for i in range(header_row + 1, len(df)):
        row = df.iloc[i]
        periodo = date_to_periodo(row.iloc[date_col])
        if not periodo:
            continue
        record = {"Periodo": periodo}
        for json_col, idx in col_indexes.items():
            record[json_col] = to_num(row.iloc[idx])
        records.append(record)
    return records


def merge_records(xlsx_records: list[dict], ppt_records: list[dict]) -> list[dict]:
    """XLSX como base (historia larga, anual preciso). PPT override para columnas mensuales."""
    by_periodo: dict[str, dict] = {r["Periodo"]: dict(r) for r in xlsx_records}

    for r in ppt_records:
        p = r["Periodo"]
        base = by_periodo.get(p, {"Periodo": p})
        for k, v in r.items():
            if k == "Periodo":
                continue
            if "Mensual" in k:
                base[k] = v
            elif k not in base or base[k] is None:
                base[k] = v
        by_periodo[p] = base

    return sorted(by_periodo.values(), key=lambda r: periodo_sort_key(r["Periodo"]))


def rebuild_json(existing: dict, records: list[dict], principal_field: str | None,
                 ppt_date: str | None) -> dict:
    periodos = [r["Periodo"] for r in records]
    updated = dict(existing)
    updated["periodos"] = periodos
    updated["series"] = records
    if ppt_date:
        updated["ultima_actualizacion"] = ppt_date

    if principal_field:
        serie_principal = [r.get(principal_field) for r in records]
        ma12 = compute_ma12(serie_principal)
        benchmarks = dict(existing.get("benchmarks") or {})
        benchmarks[f"ma12_{principal_field}"] = ma12
        updated["benchmarks"] = benchmarks

    return updated


def ingest_indicator(presentation: Presentation, xlsx_path: Path | None,
                     indicator_id: str, cfg: dict, dry_run: bool = False) -> dict:
    ppt_cfg = cfg.get("ppt_table")
    xlsx_cfg = cfg.get("embedded_xlsx")
    principal_field = cfg.get("principal_field")

    ppt_records: list[dict] = []
    xlsx_records: list[dict] = []
    sources = []

    if ppt_cfg:
        ppt_records = extract_ppt_records(presentation, ppt_cfg)
        sources.append(f"ppt({len(ppt_records)} per)")

    if xlsx_cfg and xlsx_path:
        xlsx_records = extract_xlsx_records(xlsx_path, xlsx_cfg)
        sources.append(f"xlsx({len(xlsx_records)} per)")
    elif xlsx_cfg and not xlsx_path:
        print(f"WARN [{indicator_id}]: embedded_xlsx configurado pero no se proporcionó --xlsx", file=sys.stderr)

    if not ppt_records and not xlsx_records:
        raise RuntimeError(f"[{indicator_id}] ninguna fuente disponible")

    merged = merge_records(xlsx_records, ppt_records)

    json_path = DATA_DIR / f"{indicator_id}.json"
    if not json_path.exists():
        raise FileNotFoundError(f"No existe {json_path}")
    existing = json.loads(json_path.read_text(encoding="utf-8"))

    ppt_date = parse_ppt_date(presentation)
    updated = rebuild_json(existing, merged, principal_field, ppt_date)

    summary = {
        "indicator": indicator_id,
        "sources": sources,
        "periodos": len(updated["periodos"]),
        "primer_periodo": updated["periodos"][0] if updated["periodos"] else None,
        "ultimo_periodo": updated["periodos"][-1] if updated["periodos"] else None,
        "ppt_date": ppt_date,
        "dry_run": dry_run,
    }

    if principal_field and updated.get("benchmarks"):
        ma12_arr = updated["benchmarks"].get(f"ma12_{principal_field}") or []
        ma12_validos = sum(1 for x in ma12_arr if x is not None)
        summary["ma12_validos"] = f"{ma12_validos}/{len(ma12_arr)}"

    if not dry_run:
        json_path.write_text(
            json.dumps(updated, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        summary["wrote"] = str(json_path)
    else:
        summary["preview_series_tail"] = updated["series"][-1] if updated["series"] else None

    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("pptx", help="Ruta al Reporte mensual .pptx")
    parser.add_argument("--xlsx", help="Ruta al XLSX de datos embebidos del reporte")
    parser.add_argument("--indicator", help="ID del indicador a ingestar. Si se omite, procesa todos.")
    parser.add_argument("--dry-run", action="store_true", help="No escribe el JSON, imprime resumen")
    args = parser.parse_args()

    pptx_path = Path(args.pptx)
    if not pptx_path.exists():
        print(f"ERROR: no existe {pptx_path}", file=sys.stderr)
        sys.exit(1)

    xlsx_path = Path(args.xlsx) if args.xlsx else None
    if xlsx_path and not xlsx_path.exists():
        print(f"ERROR: no existe {xlsx_path}", file=sys.stderr)
        sys.exit(1)

    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    mapping = config["indicadores"]

    if args.indicator:
        if args.indicator not in mapping:
            print(f"ERROR: indicator '{args.indicator}' no está en {CONFIG_PATH}", file=sys.stderr)
            sys.exit(1)
        targets = [args.indicator]
    else:
        targets = list(mapping.keys())

    presentation = Presentation(str(pptx_path))

    errors = []
    for indicator_id in targets:
        try:
            summary = ingest_indicator(presentation, xlsx_path, indicator_id,
                                       mapping[indicator_id], args.dry_run)
            print(json.dumps(summary, ensure_ascii=False, indent=2))
        except Exception as e:
            errors.append((indicator_id, str(e)))
            print(f"ERROR [{indicator_id}]: {e}", file=sys.stderr)

    if errors:
        sys.exit(2)


if __name__ == "__main__":
    main()
